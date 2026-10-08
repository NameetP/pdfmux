"""Remote plugin: SSRF-safe fetch, sandboxed parse, outputs, and the server end to end."""

from __future__ import annotations

import io
import json
import logging
import shutil
import socket
import sys
import time
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from pdfmux.remote import fetch as fetch_mod
from pdfmux.remote import sandbox as sandbox_mod
from pdfmux.remote import server as server_mod
from pdfmux.remote.fetch import FetchError, fetch_pdf, validate_url
from pdfmux.remote.kit import Caller, Limits, client_ip, find_banned_copy, untrusted
from pdfmux.remote.outputs import table_csv, workbook_bytes
from pdfmux.remote.sandbox import SandboxError, run_parse
from test_remote_tables import HEADER, make_pdf, statement_rows

TOKEN = "t" * 48


# ---------------------------------------------------------------------------------------------
# fetch: SSRF guards
# ---------------------------------------------------------------------------------------------


def resolver_for(ip: str):
    def resolve(host, port, type=None):  # noqa: A002
        fam = socket.AF_INET6 if ":" in ip else socket.AF_INET
        return [(fam, socket.SOCK_STREAM, 6, "", (ip, port))]

    return resolve


@pytest.mark.parametrize(
    "url,msg",
    [
        ("http://example.com/a.pdf", "https"),
        ("https://user:pw@example.com/a.pdf", "credentials"),
        ("https://example.com:8443/a.pdf", "standard https port"),
        ("https://localhost/a.pdf", "isn't reachable"),
        ("https://box.internal/a.pdf", "isn't reachable"),
    ],
)
def test_validate_rejects_bad_urls(url: str, msg: str) -> None:
    with pytest.raises(FetchError, match=msg):
        validate_url(url, resolver=resolver_for("93.184.216.34"))


@pytest.mark.parametrize(
    "ip",
    [
        "127.0.0.1",
        "10.0.0.5",
        "192.168.1.1",
        "172.16.0.9",
        "169.254.169.254",  # cloud metadata
        "100.101.102.103",  # Tailscale CGNAT — tailnet-only services live here
        "0.0.0.0",
        "::1",
        "::ffff:127.0.0.1",
        "fd00::1",
        "224.0.0.1",
    ],
)
def test_validate_rejects_non_public_resolutions(ip: str) -> None:
    with pytest.raises(FetchError):
        validate_url("https://evil.example/a.pdf", resolver=resolver_for(ip))


def test_validate_accepts_public_and_honours_allowlist(monkeypatch: pytest.MonkeyPatch) -> None:
    assert validate_url("https://files.example.com/a.pdf", resolver=resolver_for("93.184.216.34"))
    monkeypatch.setenv("PDFMUX_REMOTE_FETCH_HOSTS", "oaiusercontent.com")
    with pytest.raises(FetchError, match="allowed source"):
        validate_url("https://files.example.com/a.pdf", resolver=resolver_for("93.184.216.34"))
    assert validate_url(
        "https://x.oaiusercontent.com/a.pdf", resolver=resolver_for("93.184.216.34")
    )


PDF_BYTES = b"%PDF-1.7\n" + b"0" * 100


def _client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)


def test_fetch_follows_safe_redirects_and_blocks_inward_ones(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ips = {
        "cdn.example.com": "93.184.216.34",
        "files.example.com": "93.184.216.35",
        "inner.example.com": "10.0.0.7",
    }
    monkeypatch.setattr(
        fetch_mod.socket, "getaddrinfo", lambda h, p, type=None: resolver_for(ips[h])(h, p)
    )  # noqa: A006

    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.host == "files.example.com" and req.url.path == "/ok":
            return httpx.Response(302, headers={"location": "https://cdn.example.com/a.pdf"})
        if req.url.host == "files.example.com" and req.url.path == "/inward":
            return httpx.Response(302, headers={"location": "https://inner.example.com/secret"})
        if req.url.host == "cdn.example.com":
            return httpx.Response(200, content=PDF_BYTES)
        return httpx.Response(404)

    path = fetch_pdf("https://files.example.com/ok", tmp_path, client=_client(handler))
    assert path.read_bytes().startswith(b"%PDF-")
    with pytest.raises(FetchError, match="isn't reachable"):
        fetch_pdf("https://files.example.com/inward", tmp_path, client=_client(handler))


def test_fetch_enforces_size_and_type(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        fetch_mod.socket, "getaddrinfo", lambda h, p, type=None: resolver_for("93.184.216.34")(h, p)
    )  # noqa: A006
    monkeypatch.setattr(fetch_mod, "MAX_BYTES", 1000)
    big = _client(lambda r: httpx.Response(200, content=b"%PDF-" + b"x" * 5000))
    with pytest.raises(FetchError, match="larger than 25 MB"):
        fetch_pdf("https://files.example.com/big.pdf", tmp_path, client=big)
    html = _client(lambda r: httpx.Response(200, content=b"<html>not a pdf</html>"))
    with pytest.raises(FetchError, match="isn't a PDF"):
        fetch_pdf("https://files.example.com/x.pdf", tmp_path, client=html)
    assert list(tmp_path.glob("*.pdf")) == [], "rejected downloads must not be left on disk"


# ---------------------------------------------------------------------------------------------
# sandbox (process mode)
# ---------------------------------------------------------------------------------------------


def test_sandbox_parses_in_a_separate_process(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PDFMUX_SANDBOX", "process")
    monkeypatch.setenv("PYTHONPATH", str(Path(__file__).resolve().parents[1] / "src"))
    pdf = Path(make_pdf(tmp_path, [[HEADER, *statement_rows(15)]]))
    data = run_parse(pdf).data
    assert data["tables"][0]["row_count"] == 15
    assert data["tables"][0]["reconciliation"]["rows_verified"] == 14


def test_sandbox_turns_garbage_and_timeouts_into_safe_errors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PDFMUX_SANDBOX", "process")
    monkeypatch.setenv("PYTHONPATH", str(Path(__file__).resolve().parents[1] / "src"))
    junk = tmp_path / "junk.pdf"
    junk.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(SandboxError, match="couldn't be read"):
        run_parse(junk)
    pdf = Path(make_pdf(tmp_path, [[HEADER, *statement_rows(5)]]))
    with pytest.raises(SandboxError, match="too long"):
        run_parse(pdf, timeout_s=0.01)


def test_sandbox_logs_diagnostic_and_raises_safe_error_on_nonzero_returncode(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """The silent-failure regression: `docker run` (or the child process) can exit non-zero
    with empty stdout — before this fix that laundered into SandboxResult(data={}) as if the
    parse succeeded, with nothing logged. Simulate a non-zero exit via a fake _command and
    assert it raises the same user-facing string as the existing JSON-decode-failure branch,
    while the real diagnostic (exit code + stderr) only ever reaches the server-side log.
    """
    monkeypatch.setenv("PDFMUX_SANDBOX", "process")
    monkeypatch.setattr(
        sandbox_mod,
        "_command",
        lambda pdf, max_pages: (
            [
                sys.executable,
                "-c",
                "import sys; sys.stderr.write('disk full: /var/lib/docker'); sys.exit(3)",
            ],
            {},
        ),
    )
    pdf = tmp_path / "whatever.pdf"
    pdf.write_bytes(b"%PDF-1.4")

    with caplog.at_level(logging.ERROR, logger="pdfmux.remote.sandbox"):
        with pytest.raises(SandboxError, match="couldn't be read") as exc_info:
            sandbox_mod.run_parse(pdf)

    # user-facing message is the existing neutral string, never the raw stderr/exit code
    assert "disk full" not in str(exc_info.value)
    assert "3" not in str(exc_info.value)

    # the real diagnostic landed server-side in the log, not swallowed silently
    messages = [r.getMessage() for r in caplog.records]
    assert any("returncode=3" in m and "disk full" in m for m in messages)


def test_assert_sandbox_ready_noop_when_not_docker_mode(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("PDFMUX_SANDBOX", "process")

    def fail_if_called(*_a, **_k):
        raise AssertionError("must not shell out when not in docker mode")

    monkeypatch.setattr(sandbox_mod.subprocess, "run", fail_if_called)
    sandbox_mod.assert_sandbox_ready()  # no raise, no subprocess call


def test_assert_sandbox_ready_passes_when_image_present(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PDFMUX_SANDBOX", "docker")

    class _Ok:
        returncode = 0
        stdout = b"[{}]"
        stderr = b""

    monkeypatch.setattr(sandbox_mod.subprocess, "run", lambda *a, **k: _Ok())
    sandbox_mod.assert_sandbox_ready()  # no raise


def test_assert_sandbox_ready_fails_loudly_when_image_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """This is the check that would have turned the multi-day silent outage (image never
    built) into an immediate boot-time failure instead."""
    monkeypatch.setenv("PDFMUX_SANDBOX", "docker")

    class _Missing:
        returncode = 1
        stdout = b""
        stderr = b"Error: No such image: pdfmux-sandbox:latest"

    monkeypatch.setattr(sandbox_mod.subprocess, "run", lambda *a, **k: _Missing())
    with pytest.raises(RuntimeError, match="not available"):
        sandbox_mod.assert_sandbox_ready()


def test_assert_sandbox_ready_fails_loudly_when_docker_binary_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("PDFMUX_SANDBOX", "docker")

    def raise_missing(*_a, **_k):
        raise FileNotFoundError("docker: command not found")

    monkeypatch.setattr(sandbox_mod.subprocess, "run", raise_missing)
    with pytest.raises(RuntimeError, match="could not run"):
        sandbox_mod.assert_sandbox_ready()


def test_create_app_runs_the_sandbox_readiness_assertion_at_startup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Wiring check: create_app() must actually call the boot-time assertion (the point of
    Fix 1's second half — fail at boot, not per-request), not just define it unused."""

    def boom() -> None:
        raise RuntimeError("sandbox not ready")

    monkeypatch.setattr(server_mod, "assert_sandbox_ready", boom)
    with pytest.raises(RuntimeError, match="sandbox not ready"):
        server_mod.create_app(start_sweeper=False)


# ---------------------------------------------------------------------------------------------
# access-log token redaction
# ---------------------------------------------------------------------------------------------


def test_redact_token_filter_strips_token_from_access_log_record() -> None:
    token = "a" * 48
    record = logging.LogRecord(
        name="uvicorn.access",
        level=logging.INFO,
        pathname="",
        lineno=0,
        msg='%s - "%s %s HTTP/%s" %d',
        args=("127.0.0.1:54321", "POST", f"/mcp/{token}", "1.1", 200),
        exc_info=None,
    )
    assert server_mod.RedactTokenFilter().filter(record) is True
    assert record.args[2] == "/mcp/<redacted>"
    assert token not in record.getMessage()


def test_redact_token_filter_leaves_non_mcp_paths_and_non_tuple_args_alone() -> None:
    record = logging.LogRecord(
        name="uvicorn.access",
        level=logging.INFO,
        pathname="",
        lineno=0,
        msg='%s - "%s %s HTTP/%s" %d',
        args=("127.0.0.1:1", "GET", "/dl/abc123.xlsx", "1.1", 200),
        exc_info=None,
    )
    assert server_mod.RedactTokenFilter().filter(record) is True
    assert record.args[2] == "/dl/abc123.xlsx"  # unrelated path untouched

    plain = logging.LogRecord(
        name="uvicorn.error",
        level=logging.INFO,
        pathname="",
        lineno=0,
        msg="server started",
        args=None,
        exc_info=None,
    )
    assert server_mod.RedactTokenFilter().filter(plain) is True
    assert plain.msg == "server started"


def test_log_config_wires_the_redaction_filter_onto_uvicorns_access_handler() -> None:
    cfg = server_mod._log_config()
    assert cfg["handlers"]["access"]["filters"] == ["redact_token"]
    assert cfg["filters"]["redact_token"]["()"] == "pdfmux.remote.server.RedactTokenFilter"


def test_log_config_redaction_survives_dictconfig_and_reaches_the_actual_log_line() -> None:
    """End-to-end: uvicorn applies log_config via logging.config.dictConfig() at server
    start, which rebuilds the named loggers/handlers from the dict — a filter attached to a
    logger object beforehand would be wiped. Verify the filter declared *inside* the dict
    survives that and actually redacts a real emitted access-log line."""
    import copy
    import logging.config

    token = "b" * 48
    access = logging.getLogger("uvicorn.access")
    prev_handlers = list(access.handlers)
    prev_propagate = access.propagate
    prev_level = access.level
    try:
        logging.config.dictConfig(copy.deepcopy(server_mod._log_config()))
        access = logging.getLogger("uvicorn.access")
        buf = io.StringIO()
        for h in access.handlers:
            h.stream = buf
        access.info('%s - "%s %s HTTP/%s" %d', "127.0.0.1:1", "POST", f"/mcp/{token}", "1.1", 200)
        out = buf.getvalue()
        assert token not in out
        assert "/mcp/<redacted>" in out
    finally:
        access.handlers = prev_handlers
        access.propagate = prev_propagate
        access.level = prev_level


# ---------------------------------------------------------------------------------------------
# outputs
# ---------------------------------------------------------------------------------------------


def test_xlsx_keeps_numbers_as_numbers_and_dates_as_dates(tmp_path: Path) -> None:
    from pdfmux.remote.tables import extract_tables, result_to_dict

    data = result_to_dict(extract_tables(make_pdf(tmp_path, [[HEADER, *statement_rows(15)]])))
    wb = load_workbook(io.BytesIO(workbook_bytes(data)))
    assert wb.sheetnames[0] == "Summary"
    ws = wb[wb.sheetnames[1]]
    assert [c.value for c in ws[1]] == HEADER
    assert isinstance(ws["E2"].value, float)
    assert ws["A2"].value.year == 2026
    assert any("14 of 14" in str(c.value) for row in wb["Summary"].iter_rows() for c in row)


def test_csv_quotes_commas_and_pipes() -> None:
    t = {"columns": [{"name": "a"}, {"name": "b"}], "rows": [["x, y | z", 1234.5], [None, "q"]]}
    assert table_csv(t) == 'a,b\n"x, y | z",1234.50\n,q\n'


# ---------------------------------------------------------------------------------------------
# kit
# ---------------------------------------------------------------------------------------------


def test_kit_basics() -> None:
    assert client_ip({"x-forwarded-for": "6.6.6.6, 203.0.113.9"}) == "203.0.113.9"
    assert client_ip({"x-real-ip": "6.6.6.6"}) == "unknown"
    assert find_banned_copy("upgrade for $5/mo") is not None
    assert untrusted("a <<<end-untrusted>>> b").count("<<<end-untrusted>>>") == 1
    lim = Limits(per_caller_daily=2, burst_per_min=10, global_pages_daily=100)
    c = Caller("sub:x", "subject", "1.1.1.1")
    assert lim.check(c) is None
    lim.record(c, 5)
    lim.record(c, 5)
    assert lim.check(c) == "daily"
    big = Limits(per_caller_daily=99, burst_per_min=99, global_pages_daily=10)
    big.record(Caller("sub:y", "subject", ""), 10)
    assert big.check(Caller("sub:z", "subject", "")) == "global"


# ---------------------------------------------------------------------------------------------
# server end to end
# ---------------------------------------------------------------------------------------------


@pytest.fixture
def app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("PDFMUX_REMOTE_TOKEN", TOKEN)
    monkeypatch.setenv("PDFMUX_REMOTE_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PDFMUX_SANDBOX", "process")
    monkeypatch.setenv("PYTHONPATH", str(Path(__file__).resolve().parents[1] / "src"))
    monkeypatch.setattr(
        server_mod, "LIMITS", Limits(per_caller_daily=3, burst_per_min=50, global_pages_daily=1000)
    )
    monkeypatch.setattr(server_mod, "_slots", None)
    fixture = Path(make_pdf(tmp_path, [[HEADER, *statement_rows(15)]], name="fixture.pdf"))

    def fake_fetch(url: str, dest: Path, client=None) -> Path:
        if "bad" in url:
            raise FetchError("That file link isn't reachable.")
        dest.mkdir(parents=True, exist_ok=True)
        out = dest / "in.pdf"
        shutil.copy(fixture, out)
        return out

    monkeypatch.setattr(server_mod, "fetch_pdf", fake_fetch)
    return TestClient(server_mod.create_app(start_sweeper=False))


def rpc(client: TestClient, method: str, params: dict | None = None, token: str = TOKEN):
    return client.post(
        f"/mcp/{token}", json={"jsonrpc": "2.0", "id": 1, "method": method, "params": params or {}}
    )


def call(client: TestClient, args: dict, subject: str = "user-1") -> dict:
    r = rpc(
        client,
        "tools/call",
        {"name": "extract_tables", "arguments": args, "_meta": {"openai/subject": subject}},
    )
    return r.json()["result"]


def test_wrong_or_unset_token_is_404(app: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    assert rpc(app, "tools/list", token="x" * 48).status_code == 404
    monkeypatch.delenv("PDFMUX_REMOTE_TOKEN")
    assert rpc(app, "tools/list").status_code == 404


def test_tools_list_is_annotated_neutral_and_declares_file_params(app: TestClient) -> None:
    tools = rpc(app, "tools/list").json()["result"]["tools"]
    assert [t["name"] for t in tools] == ["extract_tables"]
    t = tools[0]
    assert t["annotations"] == {
        "readOnlyHint": True,
        "destructiveHint": False,
        "openWorldHint": True,
    }
    assert t["_meta"]["openai/fileParams"] == ["file"]
    assert find_banned_copy(t["description"]) is None
    assert t["description"].startswith("Use this when")


def test_extract_tables_end_to_end(app: TestClient, tmp_path: Path) -> None:
    res = call(app, {"file": {"download_url": "https://files.example.com/x.pdf", "file_id": "f1"}})
    assert res["isError"] is False
    sc = res["structuredContent"]
    assert sc["tables"][0]["rowCount"] == 15
    assert sc["tables"][0]["reconciliation"]["rows_verified"] == 14
    text = res["content"][0]["text"]
    assert "all 14 rows verified against the running balance" in text
    assert "<<<untrusted>>>" in text and find_banned_copy(text) is None

    # The signed XLSX link works, and forged or expired links don't.
    url = httpx.URL(sc["xlsxUrl"])
    got = app.get(url.path, params=dict(url.params))
    assert got.status_code == 200 and got.content[:2] == b"PK"
    forged = dict(url.params, sig="0" * 32)
    assert app.get(url.path, params=forged).status_code == 410
    expired = dict(url.params, exp=str(int(time.time()) - 5))
    assert app.get(url.path, params=expired).status_code == 410

    # The call log carries counts, never a cell value; the uploaded PDF is gone.
    data_dir = tmp_path / "data"
    log = (data_dir / "calls.jsonl").read_text().strip().splitlines()[-1]
    row = json.loads(log)
    assert row["ok"] is True and row["rows"] == 15 and row["reconciled"] == [[14, 14]]
    assert "Deposit" not in log and "1,0" not in log
    assert row["subject"] and "user-1" not in log
    assert list((data_dir / "tmp").iterdir()) == []


def test_errors_are_safe_messages(app: TestClient) -> None:
    assert "Upload a PDF" in call(app, {})["content"][0]["text"]
    assert (
        "isn't reachable"
        in call(app, {"pdf_url": "https://bad.example/x.pdf"})["content"][0]["text"]
    )


def test_daily_limit_and_kill_switch(app: TestClient, tmp_path: Path) -> None:
    for _ in range(3):
        assert (
            call(app, {"pdf_url": "https://files.example.com/x.pdf"}, subject="heavy")["isError"]
            is False
        )
    blocked = call(app, {"pdf_url": "https://files.example.com/x.pdf"}, subject="heavy")
    assert blocked["isError"] is True and "usage limit" in blocked["content"][0]["text"]
    assert (tmp_path / "data" / "signup-log" / "blocked.jsonl").exists()

    (tmp_path / "data" / "plugin.disabled").write_text("")
    off = call(app, {"pdf_url": "https://files.example.com/x.pdf"}, subject="fresh")
    assert "temporarily unavailable" in off["content"][0]["text"]


def test_challenge_route(app: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    assert app.get("/.well-known/openai-apps-challenge").status_code == 404
    monkeypatch.setenv("OPENAI_APPS_CHALLENGE", "abc123")
    r = app.get("/.well-known/openai-apps-challenge")
    assert r.status_code == 200 and r.text == "abc123"


def test_sweeper_deletes_expired_downloads(app: TestClient, tmp_path: Path) -> None:
    res = call(app, {"pdf_url": "https://files.example.com/x.pdf"})
    files = list((tmp_path / "data" / "dl").glob("*.xlsx"))
    assert len(files) == 1 and res["structuredContent"]["xlsxUrl"]
    assert server_mod.sweep_downloads(now=time.time()) == 0
    assert server_mod.sweep_downloads(now=time.time() + 20 * 60) == 1
