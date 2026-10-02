"""CSV and XLSX writers for the remote extract_tables tool.

Both are built from the typed table values, never by re-parsing markdown, so a ``|`` or comma
inside a cell can't shift columns, and numbers land in Excel as numbers rather than text.
"""

from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SHEET_NAME_MAX = 31


def table_csv(table: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([c["name"] for c in table["columns"]])
    for row in table["rows"]:
        w.writerow(["" if v is None else _csv_value(v) for v in row])
    return buf.getvalue()


def _csv_value(v: Any) -> str:
    if isinstance(v, float):
        return f"{v:.2f}" if abs(v - round(v, 2)) < 1e-9 else repr(v)
    return str(v)


def _sheet_title(i: int, t: dict[str, Any]) -> str:
    pages = (
        f"p{t['page_start']}"
        if t["page_start"] == t["page_end"]
        else f"p{t['page_start']}-{t['page_end']}"
    )
    return f"Table {i + 1} ({pages})"[:SHEET_NAME_MAX]


def workbook_bytes(result: dict[str, Any], source_name: str = "document.pdf") -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Source", source_name])
    summary.append(["Pages read", result["pages"]])
    if result.get("truncated"):
        summary.append(["Note", "Only the first pages were read (page limit)."])
    if result.get("scanned_pages"):
        summary.append(
            [
                "Scanned pages (no text layer, not extracted)",
                ", ".join(map(str, result["scanned_pages"])),
            ]
        )
    summary.append([])
    summary.append(["Table", "Pages", "Rows", "Rows verified against running balance"])
    for c in summary[summary.max_row]:
        c.font = Font(bold=True)

    for i, t in enumerate(result["tables"]):
        rec = t.get("reconciliation")
        verified = (
            f"{rec['rows_verified']} of {rec['rows_checked']}" if rec else "n/a (no balance column)"
        )
        title = _sheet_title(i, t)
        summary.append([title, f"{t['page_start']}-{t['page_end']}", t["row_count"], verified])

        ws = wb.create_sheet(title)
        ws.append([c["name"] for c in t["columns"]])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        types = [c["type"] for c in t["columns"]]
        for row in t["rows"]:
            out: list[Any] = []
            for v, ty in zip(row, types):
                if ty == "date" and isinstance(v, str) and len(v) == 10 and v[4] == "-":
                    try:
                        out.append(date.fromisoformat(v))
                        continue
                    except ValueError:
                        pass
                out.append(v)
            ws.append(out)
        for col_idx, ty in enumerate(types, start=1):
            letter = get_column_letter(col_idx)
            for cell in ws[letter][1:]:
                if ty == "number" and isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif ty == "date" and isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
            width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=8)
            ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
        ws.freeze_panes = "A2"

    for col in ("A", "B", "C", "D"):
        summary.column_dimensions[col].width = 34 if col in ("A", "D") else 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
